"""Parse the Çöme Ahmet family tree from the visual ODS layout.

Source of truth: aile-soyu-*.ods in this directory (Mehmet Bey's spreadsheet).
Outputs (both written next to this script):
  - people.json — long-form, human-readable schema
  - data.js     — compact schema consumed by the frontend (window.PEOPLE_DATA)

Usage:
    pip install openpyxl odfpy
    python build_tree.py
"""

import json
import re
import sys
from pathlib import Path

# --------------------------------------------------------------------------
# Paths (relative to this script — runnable on any machine)
# --------------------------------------------------------------------------
HERE = Path(__file__).resolve().parent

# Pick the most recently modified .ods (or .xlsx) in the project folder.
# Lets Mehmet Bey drop a new dated copy without renaming.
def find_source_file():
    candidates = list(HERE.glob('aile-*.ods')) + list(HERE.glob('aile-*.xlsx'))
    if not candidates:
        raise FileNotFoundError(
            f"No 'aile-*.ods' or 'aile-*.xlsx' file found in {HERE}"
        )
    # Newest by mtime — most recent edit wins.
    return max(candidates, key=lambda p: p.stat().st_mtime)

SRC_PATH = find_source_file()
PEOPLE_JSON_PATH = HERE / 'people.json'
DATA_JS_PATH = HERE / 'data.js'

# --------------------------------------------------------------------------
# Layout convention (ODS file as of 2026-05)
# --------------------------------------------------------------------------
# Root: AHMET – HATİCE displayed at (row 2, col 12), but root persons are
# pre-added programmatically with stable IDs (p_ahmet_root, p_hatice_root).
#
# 5 branches × 5 columns each, with a 1-column SEPARATOR between branches:
#
#   İSMAİL    sep   İBRAHİM   sep   HACER     sep   HÜSNİYE   sep   FATMA
#   2 3 4 5 6  7   8 9 10 11 12 13  14...18  19   20...24   25   26...30
#
# (Branch order matches the ODS — siblings either way; user confirmed sequence
#  is not load-bearing.)
#
# Daughter-branch convention (Hacer, Hüsniye, Fatma):
#   Their husbands (married-in damats) are placed in the LEFT-adjacent
#   separator column, instead of in the same column under the daughter.
#   We absorb every person in a daughter's sep-col into her clump as a spouse.
# --------------------------------------------------------------------------

BRANCHES = [
    {'name': 'İSMAİL',  'cols': [2, 3, 4, 5, 6],         'sep': None, 'gender': 'M'},
    {'name': 'İBRAHİM', 'cols': [8, 9, 10, 11, 12],      'sep': 7,    'gender': 'M'},
    {'name': 'HACER',   'cols': [14, 15, 16, 17, 18],    'sep': 13,   'gender': 'F'},
    {'name': 'HÜSNİYE', 'cols': [20, 21, 22, 23, 24],    'sep': 19,   'gender': 'F'},
    {'name': 'FATMA',   'cols': [26, 27, 28, 29, 30],    'sep': 25,   'gender': 'F'},
]

# --------------------------------------------------------------------------
# Reading: support both .ods (odfpy) and .xlsx (openpyxl)
# --------------------------------------------------------------------------

def read_ods_cells(path):
    """Read every non-empty cell from an ODS file. Returns {(row, col): text}."""
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell
    from odf.text import P

    doc = load(str(path))
    table = doc.spreadsheet.getElementsByType(Table)[0]

    def cell_text(cell):
        parts = []
        for ptag in cell.getElementsByType(P):
            parts.append(''.join(
                node.data if hasattr(node, 'data')
                else (node.firstChild.data if node.firstChild and hasattr(node.firstChild, 'data') else '')
                for node in ptag.childNodes
            ))
        return '\n'.join(parts).strip()

    cells = {}
    r_idx = 0
    for row in table.getElementsByType(TableRow):
        rep_r = int(row.getAttribute('numberrowsrepeated') or 1)
        row_cells = []
        for cell in row.getElementsByType(TableCell):
            rep_c = int(cell.getAttribute('numbercolumnsrepeated') or 1)
            txt = cell_text(cell)
            for _ in range(rep_c):
                row_cells.append(txt)
        for _ in range(rep_r):
            r_idx += 1
            for c_minus_1, txt in enumerate(row_cells):
                if txt:
                    cells[(r_idx, c_minus_1 + 1)] = txt
    return cells


def read_xlsx_cells(path):
    """Read every non-empty cell from an XLSX file. Returns {(row, col): text}."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)
    ws = wb.active
    cells = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip():
                cells[(r, c)] = str(v).strip()
    return cells


def read_cells(path):
    suffix = path.suffix.lower()
    if suffix == '.ods':
        return read_ods_cells(path)
    if suffix == '.xlsx':
        return read_xlsx_cells(path)
    raise ValueError(f"Unsupported source file extension: {suffix}")


# --------------------------------------------------------------------------
# Typo / data fixes (applied AFTER reading, BEFORE parsing)
# --------------------------------------------------------------------------
# Each rule is `(matcher_predicate, replacement_text)`. Predicates check the
# raw cell string; matched cells get replaced. Keep this list short and
# documented — every entry should record WHY the fix is correct.

CELL_FIXES = [
    # 'HACER İNCEĞİZ)' is missing the opening paren before the surname.
    # Should be 'HACER (İNCEĞİZ)' — Hacer, who married into the İNCEĞİZ family.
    # Confirmed with user.
    (lambda s: s == 'HACER İNCEĞİZ)', 'HACER (İNCEĞİZ)'),
]


def apply_cell_fixes(cells):
    fixed = {}
    for k, v in cells.items():
        new_v = v
        for pred, replacement in CELL_FIXES:
            if pred(v):
                new_v = replacement
                break
        fixed[k] = new_v
    return fixed


# --------------------------------------------------------------------------
# Person/cell parsing (unchanged from the original build_tree.py)
# --------------------------------------------------------------------------

def parse_one(text):
    text = text.strip()
    notes_extras = []
    while True:
        m = re.match(r'^(.+?)\s*\(\s*([Şş])\s*\)\s*$', text)
        if m:
            text = m.group(1).strip()
            notes_extras.append('Şehit')
            continue
        break
    m = re.match(r'^(.+?)\s*\(([^)]*)\)\s*(.*)$', text)
    if m:
        name = m.group(1).strip()
        info = m.group(2).strip()
        rest = m.group(3).strip()
        bd = re.match(r'^(\d{4})\s*-\s*(\d{4})?$', info)
        if bd:
            return {'name': name, 'birth': bd.group(1), 'death': bd.group(2) or '',
                    'notes': ' '.join(notes_extras + ([rest] if rest else []))}
        bd2 = re.match(r'^(\d{4})\s*-?\s*$', info)
        if bd2:
            return {'name': name, 'birth': bd2.group(1), 'death': '',
                    'notes': ' '.join(notes_extras + ([rest] if rest else []))}
        bd3 = re.match(r'^ÖT[\s-]*(\d{4})$', info)
        if bd3:
            note_parts = ['Öldü: ' + bd3.group(1)] + notes_extras + ([rest] if rest else [])
            return {'name': name, 'birth': '', 'death': bd3.group(1),
                    'notes': ' '.join(note_parts)}
        note_parts = [info] + notes_extras + ([rest] if rest else [])
        return {'name': name, 'birth': '', 'death': '',
                'notes': ' '.join(note_parts).strip()}
    return {'name': text, 'birth': '', 'death': '', 'notes': ' '.join(notes_extras)}


def parse_cell(text):
    raw = text
    text_clean = re.sub(r'\s+', ' ', text).strip()
    if '–' in text_clean and '(' not in text_clean:
        parts = [p.strip() for p in text_clean.split('–')]
        if len(parts) == 2 and all(parts):
            return [parse_one(p) for p in parts]
    m = re.match(r'^(.+?\(\s*[^)]*\s*\))\s+(.+?\(\s*[^)]*\s*\))\s*$', text_clean)
    if m:
        return [parse_one(m.group(1)), parse_one(m.group(2))]
    if re.search(r'\)\s{2,}\S', raw) or '\t' in raw:
        m = re.match(r'^(.+?\([^)]*\))\s{2,}(.+?)$', raw.strip())
        if m:
            return [parse_one(m.group(1).strip()), parse_one(m.group(2).strip())]
    return [parse_one(text_clean)]


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def get_surname(name):
    parts = (name or '').strip().split()
    return parts[-1].upper() if parts else ''


def normalize_match(s):
    """Fold Turkish chars for case-insensitive matching."""
    s = (s or '').upper()
    for k, v in {'İ': 'I', 'Ş': 'S', 'Ğ': 'G', 'Ü': 'U', 'Ö': 'O', 'Ç': 'C'}.items():
        s = s.replace(k, v)
    return s


def find_branch(col):
    for b in BRANCHES:
        if col in b['cols']:
            return b
    return None


# --------------------------------------------------------------------------
# Main pipeline
# --------------------------------------------------------------------------

def main():
    print(f"Reading: {SRC_PATH.name}")
    cells = read_cells(SRC_PATH)
    cells = apply_cell_fixes(cells)
    print(f"  {len(cells)} non-empty cells")

    # ----- Step 1: Pull sep-col cells aside (daughter-branch spouses) -----
    # For each daughter branch, every non-empty cell in its sep column becomes
    # an additional spouse of the branch head. We collect them now (sorted by
    # row, top to bottom) and remove them from `cells` so the normal clump
    # builder doesn't make them their own clumps.
    daughter_extra_spouses = {}  # head_col -> [parsed_person_dict, ...]
    sep_cols_to_strip = set()
    for b in BRANCHES:
        if b['gender'] != 'F' or b['sep'] is None:
            continue
        sep_col = b['sep']
        sep_cols_to_strip.add(sep_col)
        head_col = b['cols'][0]
        sep_persons = []
        for (r, c) in sorted([(r, c) for (r, c) in cells if c == sep_col]):
            sep_persons.extend(parse_cell(cells[(r, c)]))
        if sep_persons:
            daughter_extra_spouses[head_col] = sep_persons
            print(f"  daughter '{b['name']}' (col {head_col}): "
                  f"{len(sep_persons)} spouse(s) absorbed from sep col {sep_col} "
                  f"({', '.join(p['name'] for p in sep_persons)})")

    # Strip all daughter sep-cols from `cells`. We don't need them as clumps
    # anymore — they've been absorbed.
    cells = {(r, c): v for (r, c), v in cells.items() if c not in sep_cols_to_strip}

    # ----- Step 2: Build clumps (consecutive non-empty rows in same col) -----
    clumps_by_col = {}
    cols_with_data = sorted({c for (r, c) in cells.keys()})
    for col in cols_with_data:
        rows_in_col = sorted([r for (r, c) in cells.keys() if c == col])
        clumps = []
        current = [rows_in_col[0]]
        for r in rows_in_col[1:]:
            if r == current[-1] + 1:
                current.append(r)
            else:
                clumps.append(current)
                current = [r]
        clumps.append(current)
        clumps_by_col[col] = clumps

    clump_objects = []
    clump_id = 0
    seen_first_clump_in_col = set()  # col -> True after we've placed its first clump
    for col in cols_with_data:
        for rows in clumps_by_col[col]:
            persons = []
            for r in rows:
                persons.extend(parse_cell(cells[(r, col)]))
            cl = {
                'id': clump_id, 'col': col,
                'first_row': rows[0], 'last_row': rows[-1], 'rows': rows,
                'persons': persons,
            }
            # If this is the FIRST clump in a daughter branch's first col,
            # append the previously-collected sep-col spouses as additional
            # spouses of the head.
            if col not in seen_first_clump_in_col:
                seen_first_clump_in_col.add(col)
                if col in daughter_extra_spouses:
                    cl['persons'] = persons + daughter_extra_spouses[col]
            clump_objects.append(cl)
            clump_id += 1

    # Sorted lookup by col -> clumps in row order
    clumps_by_col_sorted = {}
    for c in clumps_by_col:
        clumps_by_col_sorted[c] = sorted(
            [cl for cl in clump_objects if cl['col'] == c],
            key=lambda x: x['first_row']
        )

    # ----- Step 3: Determine parent of each clump -----
    def find_parent_clump(clump):
        c = clump['col']
        r = clump['first_row']
        branch = find_branch(c)
        if branch is None:
            return None  # cell outside any known branch column — orphan
        branch_start_col = branch['cols'][0]
        if c == branch_start_col:
            # First clump in branch's first col = direct child of root.
            # All later clumps in the same col descend from that first one.
            first_clump_in_col = clumps_by_col_sorted[c][0]
            if clump['id'] == first_clump_in_col['id']:
                return 'ROOT'
            return first_clump_in_col
        # Within a branch: parent = latest clump in c-1 with first_row <= r.
        prev_col_clumps = clumps_by_col_sorted.get(c - 1, [])
        candidates = [pc for pc in prev_col_clumps if pc['first_row'] <= r]
        if candidates:
            return max(candidates, key=lambda x: x['first_row'])
        return None

    for cl in clump_objects:
        parent = find_parent_clump(cl)
        if parent == 'ROOT':
            cl['parent_id'] = 'ROOT'
        elif parent is not None:
            cl['parent_id'] = parent['id']
        else:
            cl['parent_id'] = None

    # ----- Step 4: Build people list with stable IDs -----
    people = []
    clump_to_people = {}

    # Root persons are pre-added with stable IDs (referenced by data.js + frontend).
    people.append({'id': 'p_ahmet_root', 'name': 'AHMET',
                   'birth': '', 'death': '', 'notes': 'Aile kurucusu',
                   'parent_ids': [], 'spouse_ids': ['p_hatice_root'], 'gender': 'M'})
    people.append({'id': 'p_hatice_root', 'name': 'HATİCE',
                   'birth': '', 'death': '', 'notes': 'Aile kurucusu',
                   'parent_ids': [], 'spouse_ids': ['p_ahmet_root'], 'gender': 'F'})
    ROOT_PERSON_IDS = ['p_ahmet_root', 'p_hatice_root']

    for cl in clump_objects:
        pids = []
        for i, p in enumerate(cl['persons']):
            pid = f"p{cl['id']}_{i}"
            person = {
                'id': pid,
                'name': p['name'],
                'birth': p['birth'],
                'death': p['death'],
                'notes': p['notes'],
                'parent_ids': [],
                'spouse_ids': [],
                'gender': '',
                '_clump_id': cl['id'],
                '_position_in_clump': i,
                '_row': cl['first_row'],
                '_col': cl['col'],
            }
            people.append(person)
            pids.append(pid)
        clump_to_people[cl['id']] = pids
        # Within a clump with 2+ persons: link as spouses (first = blood, rest = spouses).
        if len(pids) >= 2:
            blood = pids[0]
            for sp in pids[1:]:
                blood_person = next(pp for pp in people if pp['id'] == blood)
                sp_person = next(pp for pp in people if pp['id'] == sp)
                blood_person['spouse_ids'].append(sp)
                sp_person['spouse_ids'].append(blood)

    # ----- Step 5: Set parent_ids (multi-marriage parent → match by surname) -----
    for cl in clump_objects:
        pids = clump_to_people[cl['id']]
        if not pids:
            continue
        parent_clump_id = cl['parent_id']
        if parent_clump_id == 'ROOT':
            parent_ids = ROOT_PERSON_IDS
        elif parent_clump_id is not None:
            parent_persons = clump_to_people[parent_clump_id]
            if len(parent_persons) >= 3:
                blood_pid = parent_persons[0]
                spouse_pids = parent_persons[1:]
                child_blood_id = pids[0]
                child_blood = next(pp for pp in people if pp['id'] == child_blood_id)
                child_surname = normalize_match(get_surname(child_blood['name']))
                matched_spouses = []
                for sid in spouse_pids:
                    sp = next(pp for pp in people if pp['id'] == sid)
                    sp_surname = normalize_match(get_surname(sp['name']))
                    if sp_surname and child_surname and sp_surname == child_surname:
                        matched_spouses.append(sid)
                if len(matched_spouses) == 1:
                    parent_ids = [blood_pid, matched_spouses[0]]
                else:
                    # Ambiguous (no match, or multiple equally matching) —
                    # only blood parent is recorded; the other parent is
                    # one of the multi-marriage spouses (frontend explains this).
                    parent_ids = [blood_pid]
            else:
                parent_ids = parent_persons[:2]
        else:
            parent_ids = []
        blood_person = next(pp for pp in people if pp['id'] == pids[0])
        blood_person['parent_ids'] = parent_ids

    # ----- Step 6: Save people.json (long-form, human-readable) -----
    PEOPLE_JSON_PATH.write_text(
        json.dumps(people, ensure_ascii=False, indent=2),
        encoding='utf-8',
    )

    # ----- Step 7: Save data.js (compact schema, what the frontend reads) -----
    # Compact rules (must match what index-template.html expects):
    #   id, name      always
    #   b, d          always (empty string OK)
    #   n             only if non-empty
    #   p, s          only if non-empty arrays
    #   gender / _*   dropped (frontend doesn't use them)
    compact = []
    for p in people:
        item = {
            'id': p['id'],
            'name': p['name'],
            'b': p.get('birth', ''),
            'd': p.get('death', ''),
        }
        n = p.get('notes')
        if n:
            item['n'] = n
        par = p.get('parent_ids') or []
        if par:
            item['p'] = par
        sp = p.get('spouse_ids') or []
        if sp:
            item['s'] = sp
        compact.append(item)
    js_text = 'window.PEOPLE_DATA = ' + json.dumps(compact, ensure_ascii=False) + ';\n'
    DATA_JS_PATH.write_text(js_text, encoding='utf-8')

    # ----- Sanity report -----
    print()
    print(f"Output:")
    print(f"  {PEOPLE_JSON_PATH.name}  ({PEOPLE_JSON_PATH.stat().st_size:,} bytes)")
    print(f"  {DATA_JS_PATH.name}      ({DATA_JS_PATH.stat().st_size:,} bytes)")
    print()
    print(f"Total clumps: {len(clump_objects)}")
    print(f"Total people: {len(people)}")

    # Show direct children of root (the 5 branch heads + their spouses)
    print()
    print("Root descendants (5 branches + spouses):")
    for p in people:
        if p['parent_ids'] == ROOT_PERSON_IDS:
            spouses = [pp['name'] for pp in people if pp['id'] in p['spouse_ids']]
            sp_str = f"   eş: {spouses}" if spouses else ''
            print(f"  • {p['name']:35s}{sp_str}")

    # Sanity: per-branch person counts
    print()
    print("Per-branch person counts:")
    branch_count = {b['name']: 0 for b in BRANCHES}
    for p in people:
        col = p.get('_col')
        if col is None:
            continue
        b = find_branch(col)
        if b:
            branch_count[b['name']] += 1
    for name, cnt in branch_count.items():
        print(f"  {name:10s}: {cnt}")


if __name__ == '__main__':
    main()
